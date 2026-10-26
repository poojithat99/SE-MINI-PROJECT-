import unittest
from io import BytesIO

from docx import Document
from pypdf import PdfWriter

from src.services.text_extractor import (
    extract_text,
    TextExtractionError,
)


class TestTextExtraction(unittest.TestCase):

    def test_extract_txt(self):
        result = extract_text(
            "notes.txt",
            b"Artificial intelligence document management.",
        )
        self.assertIn("Artificial intelligence", result)

    def test_extract_docx(self):
        buffer = BytesIO()
        document = Document()
        document.add_paragraph("Document processing test.")
        document.save(buffer)

        result = extract_text("notes.docx", buffer.getvalue())
        self.assertIn("Document processing test.", result)

    def test_docx_table_cells_are_extracted(self):
        buffer = BytesIO()
        document = Document()
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Text stored in a DOCX table"
        document.save(buffer)
        self.assertIn(
            "Text stored in a DOCX table",
            extract_text("tables.docx", buffer.getvalue()),
        )

    def test_password_protected_pdf_has_readable_failure(self):
        output = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.encrypt("a-password")
        writer.write(output)
        with self.assertRaisesRegex(TextExtractionError, "password-protected"):
            extract_text("locked.pdf", output.getvalue())

    def test_corrupt_pdf_has_readable_failure(self):
        with self.assertRaisesRegex(TextExtractionError, "could not be read"):
            extract_text("broken.pdf", b"%PDF-1.4\ncorrupted")

    def test_utf8_bom_txt_is_decoded(self):
        self.assertEqual(
            extract_text("unicode.txt", b"\xef\xbb\xbfUnicode works"),
            "Unicode works",
        )

    def test_unsupported_extension_rejected(self):
        with self.assertRaises(TextExtractionError):
            extract_text("notes.exe", b"sample")

    def test_empty_document_rejected(self):
        with self.assertRaises(TextExtractionError):
            extract_text("notes.txt", b"")

    def test_pdf_without_text_fails_clearly(self):
        buffer = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.write(buffer)

        with self.assertRaisesRegex(
            TextExtractionError, "No extractable text"
        ):
            extract_text("blank.pdf", buffer.getvalue())


if __name__ == "__main__":
    unittest.main(verbosity=2)